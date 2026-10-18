"""Build the editable pitch deck and submission PDF from the same layout.

python submission/build_slides.py
Dependencies: submission/requirements-slides.txt. No inference or AWS calls.
"""
from pathlib import Path
import json
from importlib.metadata import version

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.utils import ImageReader
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'submission'
W, H = 960, 540
APP = 'https://rehearsal-act3.aa0brkr65f1ep.us-east-1.cs.amazonlightsail.com/'
REPO = 'https://github.com/AbhishekSharma55/amd-act-3'
VIDEO = 'https://youtu.be/XCLzMpVXXHU'
SOURCE = REPO + '/blob/main/'
BG, INK, MUTED = 'F6F6EF', '27332C', '647064'
ORANGE, PANEL, LINE = 'D4532B', 'E8ECE2', 'D4DBCF'
WHITE, LAVENDER, GREEN = 'FFFFFF', 'C8CAE6', '3D684E'
FONT_DIR = Path('/usr/share/fonts/truetype/liberation')
pdfmetrics.registerFont(TTFont('Body', str(FONT_DIR/'LiberationSans-Regular.ttf')))
pdfmetrics.registerFont(TTFont('Bold', str(FONT_DIR/'LiberationSans-Bold.ttf')))
pdfmetrics.registerFont(TTFont('Mono', str(FONT_DIR/'LiberationMono-Regular.ttf')))

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W/72), Inches(H/72)
prs.core_properties.title = 'Rehearsal | Forge Ahead | AMD ACT III'
prs.core_properties.subject = 'Test a factory recovery plan before changing the shift'
prs.core_properties.author = 'Abhishek Sharma and Muskan Verma'
pdf = canvas.Canvas(str(OUT/'Rehearsal_Pitch_Deck.pdf'), pagesize=(W,H))
pdf.setTitle(prs.core_properties.title)
pdf.setAuthor(prs.core_properties.author)
layout_checks = []
slide = None
page = 0


def rgb(c): return RGBColor.from_string(c)
def color(c): return tuple(int(c[i:i+2],16)/255 for i in (0,2,4))


def rect(x,y,w,h,fill=PANEL,stroke=None,r=0):
    assert x>=0 and y>=0 and x+w<=W+.1 and y+h<=H+.1,(page,x,y,w,h)
    typ=MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE
    s=slide.shapes.add_shape(typ,Inches(x/72),Inches(y/72),Inches(w/72),Inches(h/72))
    s.fill.solid();s.fill.fore_color.rgb=rgb(fill)
    if stroke:s.line.color.rgb=rgb(stroke);s.line.width=Pt(.7)
    else:s.line.fill.background()
    pdf.setFillColorRGB(*color(fill))
    if stroke:pdf.setStrokeColorRGB(*color(stroke));pdf.setLineWidth(.7)
    if r:pdf.roundRect(x,H-y-h,w,h,r,fill=1,stroke=bool(stroke))
    else:pdf.rect(x,H-y-h,w,h,fill=1,stroke=bool(stroke))
    return s


def wrap(text,w,size,bold=False,mono=False):
    font='Mono' if mono else 'Bold' if bold else 'Body'
    result=[]
    for paragraph in text.split('\n'):
        line=''
        for word in paragraph.split():
            trial=f'{line} {word}'.strip()
            if line and pdfmetrics.stringWidth(trial,font,size)>w-4:
                result.append(line);line=word
            else:line=trial
        result.append(line)
    return result


def text(text,x,y,w,size=18,fill=INK,bold=False,mono=False,leading=None,link=None):
    leading=leading or size*1.17
    lines=wrap(text,w,size,bold,mono)
    height=leading*len(lines)+3
    assert y+height<=H-12,(page,text,y,height)
    for line in lines:
        assert pdfmetrics.stringWidth(line,'Mono' if mono else 'Bold' if bold else 'Body',size)<=w+.5,(page,line)
    shape=slide.shapes.add_textbox(Inches(x/72),Inches(y/72),Inches(w/72),Inches(height/72))
    tf=shape.text_frame
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    tf.word_wrap=False;tf.vertical_anchor=MSO_ANCHOR.TOP
    for i,line in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.space_before=Pt(0);p.space_after=Pt(0);p.line_spacing=Pt(leading)
        run=p.add_run();run.text=line
        run.font.name='Liberation Mono' if mono else 'Arial'
        run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=rgb(fill)
        if link:run.hyperlink.address=link
    pdf.setFillColorRGB(*color(fill));pdf.setFont('Mono' if mono else 'Bold' if bold else 'Body',size)
    for i,line in enumerate(lines):pdf.drawString(x,H-y-size-i*leading,line)
    if link:pdf.linkURL(link,(x,H-y-height,x+w,H-y),relative=0,thickness=0)
    layout_checks.append({'slide':page,'text':text,'x':x,'y':y,'w':w,'h':height})
    return height


def image(path,x,y,w,h):
    iw,ih=Image.open(path).size
    scale=min(w/iw,h/ih);dw,dh=iw*scale,ih*scale
    xx,yy=x+(w-dw)/2,y+(h-dh)/2
    slide.shapes.add_picture(str(path),Inches(xx/72),Inches(yy/72),width=Inches(dw/72),height=Inches(dh/72))
    pdf.drawImage(ImageReader(str(path)),xx,H-yy-dh,width=dw,height=dh)


def begin(kicker,dark=False,source=None):
    global slide,page
    if page:pdf.showPage()
    page+=1;slide=prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid();slide.background.fill.fore_color.rgb=rgb(INK if dark else BG)
    pdf.setFillColorRGB(*color(INK if dark else BG));pdf.rect(0,0,W,H,stroke=0,fill=1)
    for i,h in enumerate([12,20,27]):rect(38+i*7,25+27-h,4,h,ORANGE)
    text('REHEARSAL',68,29,180,12,WHITE if dark else INK,True)
    text(kicker.upper(),420,30,490,10,'CFD6C9' if dark else MUTED)
    text('FORGE AHEAD  /  AMD DEVELOPER HACKATHON: ACT III',38,513,600,8,'BFC8B9' if dark else MUTED)
    text(f'{page:02}',894,511,30,10,'BFC8B9' if dark else MUTED)
    if source:text('Evidence / source',38,490,800,8,'BFC8B9' if dark else MUTED,link=SOURCE+source)


def title(s,sub=None):
    text(s,38,83,884,37,INK,True,leading=40)
    if sub:text(sub,40,173,870,15,MUTED)


data=json.loads((ROOT/'deploy/lightsail/verification/public_live.json').read_text())
alt=data['solve']['alternatives'][0]
replay=data['replay']
j4=next(j for j in replay['jobs'] if j['job']=='J4')

# 01 — Product and real application screenshot.
begin('Intelligent Industry',True)
text('Rehearsal.',38,107,390,58,WHITE,True)
text('Test the recovery\nbefore changing\nthe shift.',40,197,365,31,WHITE,True,leading=36)
text('Checked recovery decisions\nfor production planners.',40,337,350,19,'CFD6C9')
text('Abhishek Sharma + Muskan Verma',40,430,440,13,'CFD6C9')
rect(438,104,484,353,WHITE,r=10)
image(ROOT/'deploy/lightsail/verification/desktop.png',446,112,468,330)
text('Working AWS app • Real AMD extraction • Synthetic demo data',448,462,470,11,'CFD6C9')

# 02 — Clear user and problem, tied to the implemented example.
begin('The decision',source='deploy/lightsail/verification/public_live.json')
title('One breakdown.\nSeven delivery promises.','The production planner has to decide what to move, what runs late, and what to approve.')
cards=[('10:00','M2 stops','Work is already in progress.'),('16:00','J7 must ship','Its hard deadline must hold.'),('A decision','Own the tradeoff','A revised plan needs a clear handoff.')]
for i,(number,label,desc) in enumerate(cards):
    x=38+i*301
    rect(x,234,282,196,PANEL,r=10)
    text(number,x+20,253,246,34,ORANGE,True)
    text(label,x+20,309,245,19,INK,True)
    text(desc,x+20,353,240,16,MUTED)
text('FIRST USER',40,451,110,10,MUTED,True)
text('A production planner in a make-to-order machine shop',158,445,745,18,INK)

# 03 — Actual request-to-constraint interpretation.
begin('From a message to a decision',source='deploy/lightsail/verification/public_live.json')
title('The note becomes\na plan you can inspect.')
rect(38,188,884,83,INK,r=9)
text('“Mill M2 is unavailable until 14:00. Order J7 still needs to ship at 16:00.\nOvertime requires approval.”',57,204,842,18,WHITE)
steps=['READ ON AMD','REVIEW','SOLVE + CHECK','REHEARSE','APPROVE + EXPORT']
for i,label in enumerate(steps):
    x=38+i*180
    rect(x,294,164,54,ORANGE if i==0 else PANEL,r=6)
    text(label,x+12,312,143,11,WHITE if i==0 else INK,True)
for i,(a,b) in enumerate([('Machine outage','M2 · 10:00–14:00'),('Hard deadline','J7 · 16:00'),('Proposed option','Overtime · approval needed')]):
    x=38+i*301
    text(a,x+3,374,282,12,MUTED,True)
    text(b,x+3,399,282,19,INK,True)
text('Unknown machines and missing required times are flagged for human review.',41,451,874,15,MUTED)

# 04 — Measurable, honest tradeoff.
begin('The checked recovery',source='deploy/lightsail/verification/public_live.json')
title('Save J7.\nShow the cost.')
rect(38,190,423,251,PANEL,r=10)
text('KEEP ORIGINAL MACHINES',58,211,380,12,MUTED,True)
text('Infeasible',57,251,380,39,ORANGE,True)
text('J7’s hard deadline cannot be met\nunder this policy.',58,315,366,20)
text('The hard date is not silently moved.',58,394,367,14,MUTED)
rect(483,190,439,251,INK,r=10)
text('RE-ROUTE TO QUALIFIED MACHINES',504,211,391,12,'CFD6C9',True)
text('J7 on time',502,251,391,39,WHITE,True)
text('J2: 190 minutes late',504,315,380,25,'F5B49A',True)
text('0 overtime minutes · Independent checker passed',504,393,383,13,'CFD6C9')
text('That late delivery is an explicit concession for a person to approve.',41,454,866,17)

# 05 — Counterexample and honest scenario definition.
begin('Rehearse before committing',source='deploy/lightsail/verification/public_live.json')
title('On paper, on time.\nIn rehearsal, exposed.')
text('J4',40,191,180,16,MUTED,True)
text(f'{j4["on_time"]} / {j4["of"]}',37,218,415,65,ORANGE,True)
text('simulated scenarios finish on time',41,301,414,19)
rect(40,343,398,24,'E5D0C1',r=0)
rect(40,343,398*j4['on_time']/j4['of'],24,LAVENDER)
text('Planned finish: 16:35  ·  Due: 17:00',40,387,430,18,INK,True)
rect(489,192,433,224,INK,r=10)
text('A CONCRETE COUNTEREXAMPLE',510,212,388,12,'CFD6C9',True)
text('Sampled future #24',509,246,380,27,WHITE,True)
text('J4 finishes at 18:01',510,296,382,23,'F5B49A',True)
text('61 minutes late. Its M1 operation waits\n70 minutes longer than planned.',510,343,386,17,WHITE)
text('500 samples · seed 7 · duration factors 0.9–1.3 · repairs 0–90 min late',40,444,880,13,MUTED)
text('Synthetic uncertainty; these counts are not measured factory reliability.',40,467,880,12,MUTED)

# 06 — Actual CSV data and approval semantics.
begin('Human approval and handoff',source='deploy/lightsail/verification/public_live.json')
title('A decision someone owns.\nA record you can keep.')
items=[('01','Accept the tradeoff','Required late-delivery and overtime concessions are explicit.'),('02','Approve a stored plan','The server checks its own schedule again before export.'),('03','Keep both exports','Dispatch CSV + a decision JSON with AMD trace and rehearsal evidence.')]
for i,(n,h,b) in enumerate(items):
    y=192+i*88
    text(n,40,y+4,48,20,ORANGE,True)
    text(h,96,y,350,20,INK,True)
    text(b,96,y+31,350,15,MUTED)
rect(482,192,440,252,INK,r=10)
text('ACTUAL DISPATCH EXPORT · EXCERPT',503,212,398,11,'CFD6C9',True)
csv_lines=data['approval']['dispatch_csv'].splitlines()
table='machine  job  op  start  end\n'
for row in csv_lines[3:7]:
    f=row.split(',');table+=f'{f[0]:7}  {f[2]:3}  {f[3]:2}  {f[4]}  {f[5]}\n'
text(table.rstrip(),503,253,394,15,WHITE,mono=True,leading=25)
text('Approval locks the decision. Identical retries return the same exports.',40,465,873,14,MUTED)

# 07 — Where AMD runs, and the trace proving it.
begin('Application of technology',source='docs/ARCHITECTURE.md')
title('AMD reads the message.\nThe checks control the decision.')
rect(38,191,393,104,ORANGE,r=9)
text('AMD TOKEN FACTORY',57,207,355,12,WHITE,True)
text('DeepSeek-V4-Flash',57,235,351,27,WHITE,True)
text('Typed proposals + exact source spans',57,273,353,13,WHITE)
rect(38,315,393,126,INK,r=9)
text('WEB CONTAINER · CPU',57,330,354,12,'CFD6C9',True)
text('React + FastAPI',57,357,350,22,WHITE,True)
text('OR-Tools CP-SAT → independent checker\n→ NumPy rehearsal → approval / exports',57,394,355,15,WHITE)
rect(451,191,471,250,PANEL,r=9)
text('VERIFIED LIVE REQUEST',473,208,423,12,MUTED,True)
for i,(a,b) in enumerate([('Provider','AMD Token Factory'),('Model','DeepSeek-V4-Flash'),('Host','developer.amd.com.cn'),('Latency / attempts','7,866 ms / 1 attempt'),('Fallback','Disabled; no fallback used')]):
    y=246+i*34
    text(a,473,y,148,13,MUTED)
    text(b,628,y,272,14,INK,True)
text('No GPU replay or GPU speedup is claimed. The model API is the AMD workload.',40,465,880,13,MUTED)

# 08 — Reproducible evidence, both provisional labels visible.
begin('Evidence and its limits',source='docs/LIVE_AMD_QA.md')
title('Small tests.\nVisible limits.')
headers=[('Extractor / labels',40,353),('Complete changes',420,175),('Wrong “OK”',625,125),('Review raised',775,145)]
for value,x,w in headers:text(value,x,191,w,13,MUTED,True)
rows=[('Rules / Muskan','13 / 23','2','9 / 12'),('AMD / Muskan','23 / 23','0','12 / 12'),('Same AMD / Abhishek','21 / 23','0','9 / 10')]
for i,row in enumerate(rows):
    y=221+i*54
    rect(38,y,884,47,PANEL if i!=1 else INK,r=5)
    for j,(value,(_,x,w)) in enumerate(zip(row,headers)):
        text(value,x+7,y+13,w-14,17,WHITE if i==1 else INK,j in (0,1))
text('23 team-written messages; they informed prompt tuning. Two provisional label sets.\nNo independent benchmark or final reconciled extraction score is claimed.',41,398,877,15,MUTED)
text('Also recorded: 7 frozen synthetic recovery scenarios, including an infeasible case.\nThe repository retains unchanged outcomes and tradeoffs, not just improvements.',41,452,877,13,MUTED)

# 09 — Useful target buyer and an honest validation path.
begin('Business value and the next pilot',source='README.md')
title('A clear first user.\nA measurable next step.')
rect(38,192,423,242,INK,r=10)
text('WHO USES IT',59,211,380,12,'CFD6C9',True)
text('Production planners',58,249,380,29,WHITE,True)
text('Small, make-to-order machine shops\nhandling disruption-driven replanning.',59,297,375,18,WHITE)
text('The aim: faster, clearer recovery decisions\nwith an auditable handoff.',59,374,375,17,'CFD6C9')
text('SHADOW PILOT',493,203,412,12,MUTED,True)
for i,(h,b) in enumerate([('Use real historical disruptions','Replay cases alongside a planner’s current workflow.'),('Measure the decision','Time to a correct approved plan; missed commitments and wrong accepted changes.'),('Start with the handoff','Validate CSV exports before connecting factory systems.')]):
    y=237+i*73
    text(h,493,y,415,19,INK,True)
    text(b,493,y+28,415,14,MUTED)
text('Current limits: synthetic data · one-day fixture · no ERP integration · no validated ROI',40,464,880,12,MUTED)

# 10 — Links that work without signing in.
begin('Forge Ahead',True)
text('Rehearse before\nyou commit.',39,107,870,58,WHITE,True,leading=66)
text('Abhishek Sharma + Muskan Verma',42,270,850,23,'CFD6C9')
text('Intelligent Industry  /  AMD Developer Hackathon: ACT III',42,313,870,16,'CFD6C9')
rect(40,374,423,66,ORANGE,r=8)
text('OPEN THE LIVE APP',60,395,384,19,WHITE,True,link=APP)
rect(484,374,438,66,'3C493E',r=8)
text('VIEW CODE + EVIDENCE',504,395,400,19,WHITE,True,link=REPO)
text('github.com/AbhishekSharma55/amd-act-3',42,462,850,16,'CFD6C9',link=REPO)
text('WATCH THE DEMO  /  youtu.be/XCLzMpVXXHU',42,488,850,13,'CFD6C9',True,link=VIDEO)

pdf.save()
notes=[
    'Rehearsal helps a production planner test a recovery before changing the shift. The screenshot is from the verified public AWS application. All factory data is synthetic.',
    'At 10:00, M2 is unavailable until 14:00 while J7 still has a hard 16:00 deadline. The planner needs a defensible decision and a usable dispatch handoff.',
    'The real AMD model turns a free-text note into typed, source-linked proposals. Validation and human review precede CPU planning. Conditional overtime is a proposed option, not prior approval.',
    'The saved live run proves keep-machines infeasibility for J7. Re-routing meets J7’s hard date while J2 finishes 190 minutes late. There is no overtime in this selected plan, and the independent checker passed.',
    'J4 is on time nominally, but on time in 239 of 500 scenarios under declared synthetic assumptions. Sample 24 finishes at 18:01, 61 minutes late. These are not calibrated factory probabilities.',
    'Approval reads the stored server plan, recalculates concessions and checks the schedule again. The CSV excerpt is from the saved public workflow. The JSON preserves trace and rehearsal evidence.',
    'AMD inference is the model API workload. OR-Tools, the independent checker and NumPy replay run on CPU. This is one observed successful request, not a latency benchmark; a separate browser request needed a retry.',
    'Both provisional label sets are shown. The same 23 team-written messages informed prompt tuning, so this is development evidence, not an independent benchmark. Recovery evaluation retains unchanged cases and tradeoffs.',
    'The intended user is a planner in a small make-to-order shop. The next validation step is a shadow pilot with real historical cases and measured decision quality/time. Adoption and ROI are unvalidated.',
    'The app, code and video links are clickable. The application is live with AMD configured and fallback disabled. The repository is public and includes setup, licenses, source and reproducible evidence. Demo video: ' + VIDEO
]
for s,note in zip(prs.slides,notes):s.notes_slide.notes_text_frame.text=note
prs.save(OUT/'Rehearsal_Pitch_Deck.pptx')
(OUT/'slides_layout.json').write_text(json.dumps({'pages':page,'elements':layout_checks},indent=2)+'\n')
(OUT/'requirements-slides.txt').write_text('\n'.join(f'{package}=={version(package)}' for package in ['python-pptx','reportlab'])+'\n')
print(json.dumps({'slides':page,'pdf':str(OUT/'Rehearsal_Pitch_Deck.pdf'),'pptx':str(OUT/'Rehearsal_Pitch_Deck.pptx')},indent=2))
